#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Generate TakeCare 系統文件書.docx (標楷體／Times New Roman, 單行間距)."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path("/Users/liuzhiting/Desktop/MyApp-main/系統文件書_TakeCare_2026-09-13.docx")

CN_FONT = "標楷體"
EN_FONT = "Times New Roman"
TITLE_SIZE = 14
BODY_SIZE = 12
COVER_TITLE = 22
COVER_SUB = 16


def set_run_font(run, size, bold=False):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = EN_FONT
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), EN_FONT)
    rFonts.set(qn("w:hAnsi"), EN_FONT)
    rFonts.set(qn("w:eastAsia"), CN_FONT)
    rFonts.set(qn("w:cs"), EN_FONT)


def set_para(p, size=BODY_SIZE, bold=False, align="left", space_after=6, first_line=False):
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "justify":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if first_line:
        p.paragraph_format.first_line_indent = Cm(0.74)
    for run in p.runs:
        set_run_font(run, size, bold)


def add_p(doc, text, size=BODY_SIZE, bold=False, align="left", space_after=6, first_line=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size, bold)
    set_para(p, size, bold, align, space_after, first_line)
    return p


def add_heading_cn(doc, text, level=1):
    size = TITLE_SIZE if level <= 2 else BODY_SIZE
    p = add_p(doc, text, size=size, bold=True, space_after=10)
    return p


def add_body(doc, text):
    return add_p(doc, text, align="justify", first_line=True, space_after=8)


def add_note(doc, text):
    return add_p(doc, text, space_after=8)


def add_fig(doc, caption, body_lines):
    add_p(doc, caption, bold=True, space_after=4)
    box = "\n".join(body_lines)
    p = doc.add_paragraph()
    run = p.add_run(box)
    set_run_font(run, 11, False)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.left_indent = Cm(0.5)
    return p


def shade_header(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = tcPr.makeelement(qn("w:shd"), {
        qn("w:fill"): "1F4E5F",
        qn("w:val"): "clear",
    })
    tcPr.append(shd)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(h)
        set_run_font(run, 11, True)
        run.font.color.rgb = RGBColor(255, 255, 255)
        set_para(p, 11, True, "center", 2)
        shade_header(cell)
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(str(val))
            set_run_font(run, 11, False)
            set_para(p, 11, False, "left", 2)
    doc.add_paragraph()
    last = doc.paragraphs[-1]
    last.paragraph_format.space_after = Pt(8)
    return table


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.5)
    sec.bottom_margin = Cm(2.5)
    sec.left_margin = Cm(2.5)
    sec.right_margin = Cm(2.5)
    style = doc.styles["Normal"]
    style.font.name = EN_FONT
    style.font.size = Pt(BODY_SIZE)
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), EN_FONT)
    rFonts.set(qn("w:hAnsi"), EN_FONT)
    rFonts.set(qn("w:eastAsia"), CN_FONT)

    # Cover
    add_p(doc, "（封面顏色依當年度系上規定製作）", size=11, align="center", space_after=24)
    add_p(doc, "系統文件書", size=COVER_TITLE, bold=True, align="center", space_after=18)
    add_p(doc, "TakeCare", size=COVER_SUB, bold=True, align="center", space_after=6)
    add_p(doc, "居家跌倒偵測與遠距照護系統", size=COVER_SUB, bold=True, align="center", space_after=28)
    add_p(doc, "專題名稱：TakeCare 居家跌倒偵測與遠距照護系統", align="center", space_after=6)
    add_p(doc, "GitHub：https://github.com/takecare-agent/MyApp.git", align="center", space_after=6)
    add_p(doc, "本機分支：vision-integration", align="center", space_after=6)
    add_p(doc, "文件日期：2026 年 9 月 13 日", align="center", space_after=28)
    add_p(doc, "組別：____________________", align="center", space_after=6)
    add_p(doc, "成員／學號：____________________", align="center", space_after=6)
    add_p(doc, "指導老師：____________________", align="center", space_after=18)
    add_note(doc, "繳交提醒：紙本 5 份自行裝訂；電子檔上傳 iClass 且含 GitHub 連結；原始碼完整上傳 GitHub 後將連結提供系辦助教。本機目前未要求 push，遠端可能尚未含本週全部變更。")

    doc.add_page_break()

    # 摘要
    add_heading_cn(doc, "摘要")
    add_body(
        doc,
        "本專題實作 TakeCare，一套給居家長輩、看護與家屬共用的遠距照護系統。核心是以電腦外接攝影機做即時跌倒偵測，並把事件、待辦、血壓、緊急求救與跨語溝通接到同一支行動應用程式。偵測端使用自訓 TensorFlow Lite 模型（v8）搭配 MediaPipe 姿勢閘門，區分蹲下、彎腰與確認摔倒；確認摔倒且持續未起滿 5 秒才升級為需處理的高嚴重度通報，蹲下與彎腰只記錄不推播。行動端為 React Native（mobile-expo），後端為 Node.js，介面支援繁體中文、英語、印尼語、越南語、菲律賓語與泰語。看護負責處理，家屬知情監看。目前已有可離線安裝的 Android Release APK；服務仍由本機後端與影像行程提供。"
    )
    add_body(
        doc,
        "關鍵詞：跌倒偵測、遠距照護、姿勢估測、TensorFlow Lite、React Native、多語介面"
    )

    # 1
    add_heading_cn(doc, "1. 緒論")
    add_heading_cn(doc, "1.1 背景介紹", 2)
    add_body(
        doc,
        "高齡者居家跌倒是常見且後果嚴重的事故。跌倒後若無人發現、長時間無法起身，可能造成骨折、壓傷、脫水甚至死亡。市售監視器多半只提供連續畫面或通用移動偵測，無法區分「蹲下撿東西」與「摔倒起不來」，也缺少看護與家屬的角色分工。外籍看護在臺灣居家照護中相當普遍，介面若只有中文，會造成看不懂警報、無法回報的落差。"
    )
    add_body(
        doc,
        "本專題因此把「可信的跌倒事件」與「可操作的照護流程」做在同一產品裡：攝影機負責看見，後端負責記錄與推播，App 負責讓對的人做對的事。"
    )

    add_heading_cn(doc, "1.2 專題的目的及重要性", 2)
    add_body(
        doc,
        "目的是做出可在真機演示的閉環系統：攝影機偵測、App 通報與結案、照護待辦、血壓與急救，並讓外籍看護能切換介面語言。重要性在於事件定義必須臨床上說得通——誤把蹲下當摔倒會造成警報疲勞，看護關掉通知；漏報未起則有安全風險。本專題明確採用 Event（事件）與 Alert（警報）兩層：低姿只留痕，確認摔倒未起才吵人。"
    )

    add_heading_cn(doc, "1.3 主要研究問題或目標", 2)
    add_body(doc, "（1）如何在僅有 fall／normal 兩類的輕量模型上，用姿勢閘門分開蹲下、彎腰與摔倒？")
    add_body(doc, "（2）如何定義高／極高，使「未起 5 秒」才通報，而不是高分蹲著 5 秒？")
    add_body(doc, "（3）看護與家屬的權限如何符合「知情不處理」？")
    add_body(doc, "（4）介面如何在六語下仍可讀、長詞不被截斷？")
    add_body(doc, "（5）如何在僅有本機伺服器的條件下，讓安卓組員用一支 APK 即可登入演示？")

    add_heading_cn(doc, "1.4 系統功能簡介", 2)
    add_table(
        doc,
        ["模組", "功能"],
        [
            ["跌倒偵測", "即時串流、姿勢閘門、蹲下／彎腰僅記、摔倒短片、未起升級"],
            ["監看", "即時畫面、活動歷程、證據（截圖／短片）、未讀紅點"],
            ["緊急", "長輩 SOS、119、急救選答、看護一鍵已處理"],
            ["待辦", "家屬指派、看護回報、日常紀錄、完成紀錄"],
            ["血壓", "手打與安卓 Health Connect（歐姆龍）；家屬總覽"],
            ["訊息", "照護圈聊天、語音訊息、口譯（語音引擎已凍結）"],
            ["帳號", "三角色、照護圈綁定、六語、頭貼裁切"],
        ],
    )

    add_heading_cn(doc, "1.5 系統使用對象", 2)
    add_body(doc, "受顧者（長輩）：大 SOS、日記與紀錄、被監看。")
    add_body(doc, "看護：待辦執行與回報、監看處理、急救選答、血壓量測。")
    add_body(doc, "家屬：知情監看、指派待辦、血壓總覽；不結案跌倒事件。")

    add_heading_cn(doc, "1.6 系統特色", 2)
    add_body(doc, "Event ≠ Alert：低姿紀錄與跌倒通報分開，避免警報疲勞。")
    add_body(doc, "角色權限分離：家屬知情不處理；看護一鍵已查看；SOS 一鍵已處理。")
    add_body(doc, "本機影像：不上傳 24 小時連續影，只留事件截圖與短片。")
    add_body(doc, "六語字典：系統文案走字典，不把 UI 當使用者產生內容去機翻。")
    add_body(doc, "安卓可裝 Release APK，內含 JS bundle，不需 Metro。")

    # 2
    add_heading_cn(doc, "2. 相關技術應用與重要文獻")
    add_body(
        doc,
        "本節比較居家監視、個人緊急應變（PERS）、雲端視覺與姿勢估測等路徑，說明本專題為何採用「本機輕量模型＋姿勢閘門＋App 照護閉環」。"
    )
    add_table(
        doc,
        ["技術／產品", "優點", "缺點", "本專題取捨"],
        [
            ["一般 IP Cam", "畫面完整、市面成熟", "無跌倒語意、警報吵", "不抄連續監視產品化"],
            ["寵物／嬰兒監視", "通知分群、有事件感", "場景與高齡照護不同", "只借「緊急與日常分開」"],
            ["雲端視覺 API", "省訓練、標籤多", "隱私差、抹掉自訓差異", "不採用"],
            ["浴室雷達（Vayyar 等）", "無影像、隱私佳", "硬體成本、非本階段", "列未來工作"],
            ["MediaPipe Pose", "骨架即時、裝置可跑", "需自訂閘門邏輯", "採用，擋蹲／彎升級"],
            ["TFLite 跌倒分類", "可離線、可口試", "v8 僅 fall／normal", "用狀態機補齊類別"],
            ["AHA／ACC 血壓指引", "居家血壓有依據", "非診斷工具", "警戒 130/80 作參考"],
            ["Whisper STT", "多語語音轉文字", "資源重、維護成本", "已接入；語音模組凍結"],
        ],
    )
    add_body(
        doc,
        "文獻與產品取向上，對齊護理呼叫系統「未起才升級」的精神，以及跌倒偵測常用的髖部下落、外框寬高比，而非單幀分類分數。血壓數值解讀參考 AHA／ACC 2025 高血壓指引之居家量測門檻，僅作提醒、不作診斷。語音辨識採用 OpenAI Whisper（安卓）與蘋果語音（iOS），此管線於 2026-09-12 經使用者確認後凍結，不再更動。"
    )

    # 3
    add_heading_cn(doc, "3. 系統概要設計")
    add_heading_cn(doc, "3.1 研究或開發方法及技術", 2)
    add_body(
        doc,
        "採三層架構：影像服務（Python、TensorFlow Lite、MediaPipe）負責即時推論；後端（Node.js、MongoDB、FCM）負責帳號、事件、待辦、推播與檔案索引；行動客戶端（React Native 0.81 bare）負責三角色介面。跌倒採狀態機：正常 → 疑似 → 確認摔倒 → 未起計時 → 高／極高推播。開發採迭代驗收：需求入總帳、改碼、真機檢查、寫入驗收細項。"
    )

    add_heading_cn(doc, "3.2 工具和技術的選擇理由", 2)
    add_body(doc, "React Native bare：需要原生血壓（Health Connect）、推播、麥克風與相機權限，Expo 託管不足以覆蓋。")
    add_body(doc, "本機 vision-runtime：口試可展示自訓模型，影像資料不出門，避免雲端 API 抹平專題差異。")
    add_body(doc, "MongoDB 加本機 evidence 目錄：結構化 metadata 與二進位檔分離，清監看影時不得誤刪聊天錄音。")
    add_body(doc, "六語自建字典：避免系統文案被當 UGC 機翻，臨床用語（蹲下≠滑倒、低姿勢紀錄）可人工校正。")

    add_heading_cn(doc, "3.3 處理流程", 2)
    add_body(doc, "跌倒主流程如下。")
    add_body(doc, "（1）外接攝影機 cam 0 取 1920×1080。")
    add_body(doc, "（2）MediaPipe 估姿勢；TFLite 輸出 fall 分數。")
    add_body(doc, "（3）髖夠低且符合躺姿才確認摔倒；持續蹲下或彎腰不得因滿 5 秒變成摔倒。")
    add_body(doc, "（4）蹲下／彎腰：寫低嚴重度紀錄與截圖，不推播。")
    add_body(doc, "（5）確認摔倒：錄短片；未起滿 5 秒升為高；持續未起另推極高波次（第 2、5、15、35 分鐘）。")
    add_body(doc, "（6）後端寫入事件並推播看護與家屬。")
    add_body(doc, "（7）看護一鍵已查看（可選說明）；家屬只看、不結案。")
    add_body(doc, "SOS 流程：長輩按住求救 → 看護與家屬收卡 → 看護一鍵已處理 → 家屬側同步消失。")

    add_heading_cn(doc, "3.4 檔案關連", 2)
    add_table(
        doc,
        ["元件", "主要路徑", "說明"],
        [
            ["行動 App", "mobile-expo/src/", "現行客戶端"],
            ["後端 API", "backend/index.js", "帳號、待辦、聊天、血壓、事件"],
            ["影像服務", "vision-runtime/vision_api_server.py", "v8 模型＋姿勢閘門"],
            ["鏡頭啟動", "vision-runtime/start-cam.sh", "鎖定 --cam 0"],
            ["語系", "mobile-expo/src/i18n/", "六語字典"],
            ["需求歷程", "Fall_Detection_Lab/docs/", "總帳、SPEC、驗收"],
            ["進度", "Fall_Detection_Lab/APP_PROGRESS.md", "當下 ⭐"],
            ["安裝包", "~/Desktop/TakeCare-demo/*.apk", "安卓 Release"],
            ["舊 Vite", "frontend/", "淘汰、勿當主力"],
        ],
    )

    add_heading_cn(doc, "3.5 系統架構圖", 2)
    add_fig(
        doc,
        "圖 1 系統架構",
        [
            "[外接攝影機 cam 0]",
            "          |",
            "          v",
            "[vision-runtime :8000] --detect--> [backend :5000] --REST / FCM--> [Android APK / iOS]",
            "          |                              |",
            "       /stream                      MongoDB + evidence/",
            "",
            "說明：手機不跑模型。Release APK 內含 JS，仍須連本機 :5000；監看再連 :8000 或經後端轉送。",
        ],
    )

    add_heading_cn(doc, "3.6 流程圖", 2)
    add_fig(
        doc,
        "圖 2 跌倒偵測與通報流程",
        [
            "取幀 → 姿勢估測 → TFLite 分數",
            "   |",
            "   +-- 蹲下／彎腰 --> 寫低紀錄 --> 結束（不推播）",
            "   +-- 確認摔倒 --> 錄短片 --> 未起計時",
            "                        |",
            "                        +-- < 5 秒起身 --> 僅事件",
            "                        +-- >= 5 秒 -----> 高嚴重度推播",
            "                        +-- 持續未起 ----> 極高波次",
            "                        +-- 看護已查看 --> 結案、停止波次",
        ],
    )

    add_heading_cn(doc, "3.7 使用案例圖", 2)
    add_fig(
        doc,
        "圖 3 使用案例（摘要）",
        [
            "長輩：按 SOS、寫日記、被監看、看自己的紀錄",
            "看護：處理 SOS、已查看跌倒、執行待辦、量血壓、看即時／活動、急救選答",
            "家屬：知情監看、指派待辦、看血壓、不結案跌倒",
            "系統：偵測跌倒、推播、寫證據、六語呈現",
        ],
    )
    add_table(
        doc,
        ["編號", "使用案例", "主要角色", "結果"],
        [
            ["UC1", "確認摔倒未起通報", "系統→看護／家屬", "高或極高推播"],
            ["UC2", "看護已查看", "看護", "事件結案，家屬同步"],
            ["UC3", "長輩 SOS", "長輩→看護", "一鍵已處理"],
            ["UC4", "指派與回報待辦", "家屬／看護", "今日清單與完成紀錄"],
            ["UC5", "同步血壓", "看護／長輩（安卓）", "家屬可看趨勢"],
            ["UC6", "切換介面語", "看護", "系統文案跟語系"],
        ],
    )

    add_heading_cn(doc, "3.8 甘特圖", 2)
    add_table(
        doc,
        ["期間", "工作項目"],
        [
            ["2026-07 以前", "模型 v7／v8 訓練、Lab 資料集與研究"],
            ["2026-08 上旬", "角色、警報分級規格、SOS、i18n 骨架"],
            ["2026-08 下旬", "姿勢閘門、語音訊息、異常結案改口"],
            ["2026-09 上旬", "監看時間軸、血壓、急救選答"],
            ["2026-09-12", "六語補齊、離線 APK、現行包說明"],
            ["2026-09-13", "留存決策、系統文件書、新帳號交接"],
        ],
    )
    add_fig(
        doc,
        "圖 4 甘特（示意）",
        [
            "Jul |████ 模型／Lab",
            "Aug |    ████████ 角色／警報／SOS／閘門／語音",
            "Sep |            ████████ 監看／血壓／六語／APK／文件",
        ],
    )

    add_heading_cn(doc, "3.9 活動圖", 2)
    add_fig(
        doc,
        "圖 5 看護處理摔倒活動",
        [
            "推播到達 → 開啟監看或活動 → 觀看短片／截圖",
            "→（可選）填寫或修改說明 → 按「已查看」",
            "→ 後端結案 → 家屬側同步 → 極高波次停止",
        ],
    )

    # 4
    add_heading_cn(doc, "4. 系統開發工具與使用環境")
    add_heading_cn(doc, "4.1 開發工具", 2)
    add_table(
        doc,
        ["項目", "版本／備註"],
        [
            ["Node.js", "20.x（腳本指定 20.20.2）"],
            ["React Native", "0.81 bare（mobile-expo）"],
            ["Python", "3.11 + vision-runtime .venv"],
            ["模型", "fall_detection_model_v8.tflite"],
            ["資料庫", "MongoDB"],
            ["推播", "Firebase Cloud Messaging"],
            ["安卓建置", "Gradle／Android SDK，assembleRelease"],
            ["iOS 建置", "Xcode（無現成 IPA）"],
            ["除錯", "Metro 8081、adb、curl"],
        ],
    )

    add_heading_cn(doc, "4.2 運行環境要求", 2)
    add_body(
        doc,
        "伺服器（開發與演示用 Mac）：macOS、外接 USB 攝影機鎖定 cam 0。埠：後端 5000、影像 8000；開發熱重載另開 Metro 8081。啟動順序為後端 → 鏡頭；確認指令為 bash ~/Desktop/MyApp-main/check-stack.sh。"
    )
    add_body(
        doc,
        "Android 客戶端：安裝 TakeCare-vision-integration.apk（約 54MB，含 assets/index.android.bundle）。與 Mac 同一熱點，伺服器填 http://<Mac-IP>:5000。不需 Metro，可拔 USB。熱點 IP 會變，以 ipconfig getifaddr en0 為準。"
    )
    add_body(
        doc,
        "iOS 客戶端：目前無 IPA。第一次須 USB 加 Xcode 安裝；之後開發 Reload 需與 Mac 同一熱點並開著 Metro。電腦離開熱點會載入手機內建舊 JS。"
    )
    add_body(
        doc,
        "測試帳號：patient@test.com、caregiver@test.com、family@test.com，密碼皆為 Test1234!。"
    )

    # 5
    add_heading_cn(doc, "5. 系統實作及實驗結果")
    add_heading_cn(doc, "5.1 系統功能的詳細描述", 2)
    add_body(
        doc,
        "帳號與照護圈：三角色註冊／登入，綁定同一長輩。設定可改姓名、電話、生日、性別與頭貼（拖移縮放裁切）。介面語六選一。"
    )
    add_body(
        doc,
        "跌倒與監看：即時 MJPEG、活動歷程、證據檔。低姿每次各記一筆不推播。確認摔倒未起滿 5 秒為高；之後依 2／5／15／35 分鐘極高波次。看護已查看即結案。家屬不處理。"
    )
    add_body(
        doc,
        "緊急：長輩大 SOS。看護收卡後一鍵已處理。急救選答含 119、CPR、哽塞 5+5、跌倒與 FAST 步驟。健康資訊卡供急救一眼閱讀。"
    )
    add_body(
        doc,
        "待辦與日常：家屬指派、看護回報。已完成離開今日、進入完成紀錄。日常紀錄類別可編。待辦頂部不放搜尋欄。"
    )
    add_body(
        doc,
        "血壓：手打與安卓 Health Connect 匯入（歐姆龍先寫入 Health Connect）。庫保留約 3 個月前與 6 個月前參考點供線圖。iPhone 不接蘋果健康。"
    )
    add_body(
        doc,
        "訊息：照護圈聊天、語音氣泡、口譯。語音 STT／Whisper／錄音管線已凍結。手打待辦標題跟讀取方語言。"
    )

    add_heading_cn(doc, "5.2 實驗數據與評估", 2)
    add_body(
        doc,
        "模型側評估資料在 Fall_Detection_Lab（混淆矩陣、thresholds_v8），屬研究資料，不在本文件重貼完整數字。產品側以真機兩台（華為安卓＋iPhone）驗收，不以三台模擬器為預設。"
    )
    add_body(
        doc,
        "2026-09-12：歐姆龍經 Health Connect 同步由使用者口述成功。六語字典掃描：非中文語系 0 筆繁中混入；相對中文 catalog 缺 key 為 0。急救卡與日誌卡已去掉單行截斷。"
    )
    add_body(
        doc,
        "限制：未做大規模臨床試驗；跌倒測試僅用緩慢躺下，禁止真摔。辨識幾何公式凍結，未再改姿勢閘門。極高波次在行程重啟後的持久化仍為已知限制。"
    )

    add_heading_cn(doc, "5.3 與其他相關技術應用的比較", 2)
    add_body(
        doc,
        "相對純監視器：有事件分級與角色，不是 24 小時盯畫面。相對雲端視覺 API：可離線推論、隱私較可控、口試能講自訓模型。相對單幀分類 App：有未起時間窗與姿勢閘門，降低蹲下誤報。相對商用 PERS：無專用硬體項鍊，改用攝影機加手機，成本較低但依賴取景與本機服務。"
    )

    add_heading_cn(doc, "5.4 遭遇的問題和挑戰", 2)
    add_body(doc, "（1）v8 只有 fall／normal，必須靠 MediaPipe 閘門。軀幹水平曾把彎腰當躺，已撤回該條件。")
    add_body(doc, "（2）iPhone 離熱點會載舊 JS，易被誤認功能消失。")
    add_body(doc, "（3）六語長詞截斷；系統文案曾被當 UGC 機翻，出現越語摻中文、泰文誤譯。")
    add_body(doc, "（4）口譯曾因模型額度 429 不穩；語音管線其後凍結。")
    add_body(doc, "（5）清檔腳本曾誤刪舊語音 WAV，其後禁止未明講即刪使用者資料。")
    add_body(doc, "（6）Mac 對安卓不能像隨身碟拖檔，改以 adb install -r 分發 APK。")
    add_body(doc, "（7）舊 frontend（Vite＋Capacitor）與 Lab 工作區影像腳本易與現行接線混淆，已文件化隔離。")

    # 6
    add_heading_cn(doc, "6. 結論及未來發展")
    add_heading_cn(doc, "6.1 主要貢獻", 2)
    add_body(
        doc,
        "本專題完成可演示的「偵測＋照護」閉環：自訓 TFLite 模型、姿勢閘門、角色權限、六語 App、可分發 Android APK。事件定義對齊居家照護——蹲下是低姿勢紀錄，不是滑倒，也不是健身房深蹲。歷程與需求分散記錄於 Fall_Detection_Lab 的進度、總帳、SPEC 與驗收檔，並於 2026-09-13 整理留存決策與本系統文件書。"
    )

    add_heading_cn(doc, "6.2 未來研究或發展建議", 2)
    add_body(doc, "建議在使用者同意後，以受控場景驗證真實跌倒與起身，而不是只靠緩慢躺下。建議補齊 iOS 血壓與正式 IPA。口譯若再啟動，應改官方翻譯 API 加術語表，不要非官方套件。")

    add_heading_cn(doc, "6.3 未來工作與預期挑戰", 2)
    add_body(doc, "（1）iOS 血壓直連與 TestFlight／IPA。")
    add_body(doc, "（2）官方翻譯 API 加術語表（口譯再解凍時）。")
    add_body(doc, "（3）極高波次在服務重啟後的持久化。")
    add_body(doc, "（4）使用者同意後的實機跌倒場景驗證。")
    add_body(doc, "（5）依當年度封面色印製五份紙本、iClass 上傳，並把完整原始碼推上 GitHub 供助教。")
    add_body(doc, "預期挑戰：臨床樣本不足、推播在背景被系統休眠、外籍用語持續校正、封面色與裝訂時程。")

    # 7
    add_heading_cn(doc, "7. 參考文獻")
    refs = [
        "American Heart Association / American College of Cardiology, 2025 High Blood Pressure Guideline.",
        "Google, MediaPipe Pose. https://developers.google.com/mediapipe",
        "Google, TensorFlow Lite. https://www.tensorflow.org/lite",
        "Meta, React Native. https://reactnative.dev",
        "OpenAI, Whisper. https://github.com/openai/whisper",
        "本專題規格：Fall_Detection_Lab/docs/SPEC_警報分級與異常事件_2026-08-05.md",
        "本專題需求總帳：Fall_Detection_Lab/docs/REQUIREMENTS_用戶對接總帳_2026-08-10.md",
        "本專題競品研究：Fall_Detection_Lab/docs/RESEARCH_警報與攝影機UX_2026-08-10.md",
        "本專題現行包：Fall_Detection_Lab/docs/CURRENT_現行這包_2026-09-12.md",
        "本專題留存決策：Fall_Detection_Lab/docs/PROJECT_留存與淘汰決策_2026-09-13.md",
        "原始碼遠端：https://github.com/takecare-agent/MyApp.git",
    ]
    for i, ref in enumerate(refs, 1):
        add_p(doc, f"[{i}] {ref}", space_after=4)

    add_heading_cn(doc, "附錄 A 文件與歷程對照")
    add_body(
        doc,
        "歷程沒有集中在單一檔，也沒有整包刪掉。主線在 Fall_Detection_Lab/APP_PROGRESS.md、需求總帳、驗收 PROCESS、CHECKPOINT 快照、SPEC 與桌面日常流程。本對話實際刪除僅 vision-runtime/_experiments/vision_api_server.py.bak_20260729。frontend、_experiments、Archive、Backup、使用者資料均保留。Archive 舊 APK 禁止再裝。"
    )

    add_heading_cn(doc, "附錄 B 新帳號交接")
    add_body(
        doc,
        "完整可貼上的開場指令見 Fall_Detection_Lab/docs/HANDOVER_新帳號開場指令_2026-09-13.md。只改 MyApp-main 與 Fall_Detection_Lab。語音、辨識幾何、真摔、刪使用者資料等硬規則維持不變。"
    )

    doc.save(OUT)
    print(f"Wrote {OUT} ({OUT.stat().st_size} bytes)")


if __name__ == "__main__":
    build()
